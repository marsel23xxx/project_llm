"""Turn a written document (Markdown from the LLM) or a result table into a downloadable file.

DOCX and PDF are rendered from the same parsed blocks, so both carry the full document with real
headings, paragraphs, lists and tables rather than raw Markdown marks.
"""

import io
import re
from dataclasses import dataclass, field
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services import chart_image

FORMATS = {
    "pdf": ("application/pdf", "PDF"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "Word"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "Excel"),
}

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*\s][^*]*?\*|`[^`]+`)")


_FENCE = re.compile(r"^\s*```\s*([\w-]*)\s*$")
CHART_WIDTH_CM = 14.0


@dataclass
class Block:
    kind: str  # "heading", "paragraph", "list", "table", "chart" or "code"
    text: str = ""
    chart: chart_image.ChartSpec | None = None
    level: int = 0
    ordered: bool = False
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


def parse_markdown(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", text=" ".join(paragraph)))
            paragraph.clear()

    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or re.fullmatch(r"\s*([-*_]\s*){3,}", line):  # blank line or horizontal rule
            flush()
        elif fence := _FENCE.match(line):
            flush()
            body = []
            i += 1
            while i < len(lines) and not _FENCE.match(lines[i]):
                body.append(lines[i])
                i += 1
            if fence.group(1).lower() == "chart":
                try:
                    blocks.append(Block("chart", chart=chart_image.parse_spec("\n".join(body))))
                except chart_image.ChartError:
                    pass  # a malformed chart is left out rather than drawn wrongly
            elif body:
                blocks.append(Block("code", text="\n".join(body)))
        elif match := _HEADING.match(line):
            flush()
            blocks.append(Block("heading", text=match.group(2).strip().strip("#").strip(), level=len(match.group(1))))
        elif _TABLE_ROW.match(line) and i + 1 < len(lines) and _TABLE_RULE.match(lines[i + 1]):
            flush()
            rows = [[c.strip() for c in _TABLE_ROW.match(line).group(1).split("|")]]
            i += 2
            while i < len(lines) and (row := _TABLE_ROW.match(lines[i])):
                rows.append([c.strip() for c in row.group(1).split("|")])
                i += 1
            width = len(rows[0])
            blocks.append(Block("table", rows=[(r + [""] * width)[:width] for r in rows]))
            continue
        elif (bullet := _BULLET.match(line)) or _NUMBERED.match(line):
            flush()
            ordered = bullet is None
            pattern = _NUMBERED if ordered else _BULLET
            items = []
            while i < len(lines) and (item := pattern.match(lines[i])):
                items.append(item.group(2 if ordered else 1).strip())
                i += 1
            blocks.append(Block("list", ordered=ordered, items=items))
            continue
        else:
            paragraph.append(line.strip())
        i += 1
    flush()
    return blocks


def split_inline(text: str) -> list[tuple[str, str]]:
    """Split text into (style, content) runs, where style is "", "bold", "italic" or "code"."""
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and len(part) > 4:
            runs.append(("bold", part[2:-2]))
        elif part.startswith("*") and len(part) > 2:
            runs.append(("italic", part[1:-1]))
        elif part.startswith("`") and len(part) > 2:
            runs.append(("code", part[1:-1]))
        else:
            runs.append(("", part))
    return runs


def plain(text: str) -> str:
    return "".join(content for _, content in split_inline(text))


def document_title(markdown: str, fallback: str = "Dokumen") -> str:
    for block in parse_markdown(markdown):
        if block.kind == "heading":
            return plain(block.text)[:120]
    return fallback


def word_count(markdown: str) -> int:
    return len(re.findall(r"\w+", markdown))


def format_count(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def slug(title: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
    return cleaned[:60] or "dokumen"


# ---------- DOCX ----------


def _add_runs(paragraph, text: str) -> None:
    for style, content in split_inline(text):
        run = paragraph.add_run(content)
        run.bold = style == "bold"
        run.italic = style == "italic"
        if style == "code":
            run.font.name = "Courier New"


def render_docx(markdown: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(3)
    section.top_margin = section.bottom_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = document.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = "Times New Roman", Pt(size), True
        style.font.color.rgb = None

    first_heading = True
    for block in parse_markdown(markdown):
        if block.kind == "heading":
            if first_heading:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(plain(block.text))
                run.bold, run.font.size = True, Pt(16)
                paragraph.paragraph_format.space_after = Pt(18)
                first_heading = False
            else:
                document.add_heading(plain(block.text), level=min(max(block.level - 1, 1), 3))
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(paragraph, block.text)
        elif block.kind == "list":
            for item in block.items:
                _add_runs(document.add_paragraph(style="List Number" if block.ordered else "List Bullet"), item)
        elif block.kind == "chart":
            png, _ = chart_image.render_png(block.chart)
            document.add_picture(io.BytesIO(png), width=Cm(CHART_WIDTH_CM))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.kind == "code":
            run = document.add_paragraph().add_run(block.text)
            run.font.name, run.font.size = "Courier New", Pt(10)
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(block.rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c).paragraphs[0]
                    cell.paragraph_format.line_spacing = 1.0
                    _add_runs(cell, value)
                    if r == 0:
                        for run in cell.runs:
                            run.bold = True
            document.add_paragraph()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

_BODY = ParagraphStyle("doc-body", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_JUSTIFY, spaceAfter=8)
_TITLE = ParagraphStyle("doc-title", fontName="Times-Bold", fontSize=16, leading=21, alignment=TA_CENTER, spaceAfter=20)
_HEADINGS = {
    1: ParagraphStyle("doc-h1", fontName="Times-Bold", fontSize=14, leading=19, spaceBefore=14, spaceAfter=8, keepWithNext=1),
    2: ParagraphStyle("doc-h2", fontName="Times-Bold", fontSize=12.5, leading=17, spaceBefore=10, spaceAfter=6, keepWithNext=1),
    3: ParagraphStyle("doc-h3", fontName="Times-BoldItalic", fontSize=12, leading=16, spaceBefore=8, spaceAfter=4, keepWithNext=1),
}
_LIST = ParagraphStyle("doc-list", parent=_BODY, alignment=TA_LEFT, leftIndent=20, bulletIndent=6, spaceAfter=4, bulletFontName="Times-Roman", bulletFontSize=12)
_CELL = ParagraphStyle("doc-cell", fontName="Times-Roman", fontSize=10.5, leading=13)


def _markup(text: str) -> str:
    out = []
    for style, content in split_inline(text):
        safe = escape(content)
        out.append(
            {"bold": f"<b>{safe}</b>", "italic": f"<i>{safe}</i>", "code": f'<font face="Courier">{safe}</font>'}.get(style, safe)
        )
    return "".join(out)


def _page_number(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 10)
    canvas.drawCentredString(A4[0] / 2, 1.3 * cm, str(doc.page))
    canvas.restoreState()


def _column_widths(rows: list[list[str]], total: float) -> list[float]:
    """Wider columns for longer text, so one wordy column does not make every row tall.

    Each column first gets room for its longest word (a narrow column would break "Kesehatan" in two);
    what is left is shared by the length of the text.
    """
    columns = range(len(rows[0]))
    texts = [[plain(row[c]) for row in rows] for c in columns]
    floor = [max(stringWidth(word, "Times-Bold", _CELL.fontSize) for text in texts[c] for word in text.split() or [""]) + 14 for c in columns]
    weights = [min(max(len(text) for text in texts[c]), 60) + 1 for c in columns]
    spare = total - sum(floor)
    if spare < 0:  # words too long to honour: plain proportional split
        return [total * w / sum(weights) for w in weights]
    return [floor[c] + spare * weights[c] / sum(weights) for c in columns]


def render_pdf(markdown: str) -> bytes:
    story = []
    first_heading = True
    for block in parse_markdown(markdown):
        if block.kind == "heading":
            if first_heading:
                story.append(Paragraph(_markup(plain(block.text)), _TITLE))
                first_heading = False
            else:
                story.append(Paragraph(_markup(plain(block.text)), _HEADINGS[min(max(block.level - 1, 1), 3)]))
        elif block.kind == "paragraph":
            story.append(Paragraph(_markup(block.text), _BODY))
        elif block.kind == "list":
            for n, item in enumerate(block.items, start=1):
                story.append(Paragraph(_markup(item), _LIST, bulletText=f"{n}." if block.ordered else "•"))
            story.append(Spacer(1, 4))
        elif block.kind == "chart":
            png, ratio = chart_image.render_png(block.chart)
            width = CHART_WIDTH_CM * cm
            # Kept together so a chart is never split from the page it belongs on.
            story += [KeepTogether([PdfImage(io.BytesIO(png), width=width, height=width * ratio)]), Spacer(1, 10)]
        elif block.kind == "code":
            story.append(Paragraph(escape(block.text).replace("\n", "<br/>"), ParagraphStyle("doc-code", fontName="Courier", fontSize=9.5, leading=12, spaceAfter=8)))
        elif block.kind == "table":
            data = [
                [Paragraph(f"<b>{_markup(v)}</b>" if r == 0 else _markup(v), _CELL) for v in row]
                for r, row in enumerate(block.rows)
            ]
            # splitInRow: a row taller than the space left is cut across the page break. Without it the
            # whole table jumps to the next page and leaves the rest of this one empty.
            table = Table(data, colWidths=_column_widths(block.rows, A4[0] - 6 * cm), repeatRows=1, splitInRow=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story += [table, Spacer(1, 10)]

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=3 * cm, rightMargin=3 * cm, topMargin=2.5 * cm, bottomMargin=2.5 * cm,
        title=document_title(markdown), author="Agentic Data Analyst",
    ).build(story or [Paragraph("-", _BODY)], onFirstPage=_page_number, onLaterPages=_page_number)  # fmt: skip
    return buffer.getvalue()


# ---------- XLSX ----------


def render_xlsx(columns: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hasil"
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for index, column in enumerate(columns, start=1):
        longest = max([len(str(column))] + [len(str(r[index - 1])) for r in rows[:200] if r[index - 1] is not None])
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(longest + 2, 8), 50)
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def table_as_markdown(columns: list[str], rows: list[list], limit: int = 60) -> str:
    def cell(value) -> str:
        return "-" if value is None else str(value).replace("|", "/")

    lines = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    lines += ["| " + " | ".join(cell(v) for v in row) + " |" for row in rows[:limit]]
    return "\n".join(lines)
