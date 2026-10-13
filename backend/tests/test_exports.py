import io
import json
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from app.api.v1.chat import build_export
from app.core.config import Settings
from app.services import exports
from app.services.agent.graph import build_agent, is_writing_request, stream_agent, strip_fence
from app.services.llm.base import LLMResult

PAPER = """# Energi Terbarukan di Indonesia

## Pendahuluan

Indonesia memiliki potensi **energi surya** yang besar karena berada di garis khatulistiwa. Pemanfaatannya *masih rendah* dibanding potensinya.

### Rumusan Masalah

1. Apa hambatan utama pemanfaatan energi surya?
2. Bagaimana strategi mengatasinya?

## Pembahasan

Hambatan utama meliputi biaya awal, jaringan listrik, dan regulasi.

- Biaya pemasangan `PLTS atap` masih tinggi
- Jaringan belum merata

| Sumber | Potensi | Terpasang |
|---|---|---|
| Surya | tinggi | rendah |
| Angin | sedang | rendah |

## Penutup

Rujukan perlu dilengkapi penulis.
"""


def _flat(reader: PdfReader) -> str:
    """PDF text with line wraps removed, so phrases can be searched for."""
    return " ".join(" ".join(page.extract_text() for page in reader.pages).split())


def test_markdown_is_parsed_into_document_blocks():
    blocks = exports.parse_markdown(PAPER)
    assert [b.kind for b in blocks] == [
        "heading", "heading", "paragraph", "heading", "list", "heading", "paragraph", "list", "table", "heading", "paragraph",
    ]  # fmt: skip
    assert blocks[4].ordered and blocks[4].items[1] == "Bagaimana strategi mengatasinya?"
    assert not blocks[7].ordered and blocks[8].rows == [["Sumber", "Potensi", "Terpasang"], ["Surya", "tinggi", "rendah"], ["Angin", "sedang", "rendah"]]
    assert exports.split_inline("teks **tebal** dan *miring* serta `kode`") == [
        ("", "teks "), ("bold", "tebal"), ("", " dan "), ("italic", "miring"), ("", " serta "), ("code", "kode"),
    ]  # fmt: skip
    assert exports.document_title(PAPER) == "Energi Terbarukan di Indonesia" and exports.slug("Energi & Surya 2026!") == "energi-surya-2026"


def test_docx_contains_the_whole_document_with_real_structure():
    document = Document(io.BytesIO(exports.render_docx(PAPER)))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    text = "\n".join(p.text for p in paragraphs)

    assert paragraphs[0].text == "Energi Terbarukan di Indonesia" and paragraphs[0].runs[0].bold
    assert [p.text for p in paragraphs if p.style.name.startswith("Heading")] == ["Pendahuluan", "Rumusan Masalah", "Pembahasan", "Penutup"]
    assert "berada di garis khatulistiwa" in text and "Jaringan belum merata" in text
    assert "**" not in text and "#" not in text and "`" not in text  # marks became formatting
    bold_runs = [r.text for p in paragraphs for r in p.runs if r.bold and p.style.name == "Normal"]
    assert "energi surya" in bold_runs
    assert [p.style.name for p in paragraphs if "hambatan utama pemanfaatan" in p.text] == ["List Number"]
    assert [c.text for c in document.tables[0].rows[1].cells] == ["Surya", "tinggi", "rendah"]
    assert document.sections[0].page_width.cm == pytest.approx(21, abs=0.1)


def test_pdf_contains_the_whole_document_without_markdown_marks():
    reader = PdfReader(io.BytesIO(exports.render_pdf(PAPER)))
    text = _flat(reader)
    for expected in ("Energi Terbarukan di Indonesia", "Rumusan Masalah", "berada di garis khatulistiwa", "PLTS atap", "Angin", "Penutup"):
        assert expected in text, expected
    assert "**" not in text and "##" not in text and "|" not in text
    assert reader.metadata.title == "Energi Terbarukan di Indonesia"


def test_xlsx_holds_the_result_table():
    sheet = load_workbook(io.BytesIO(exports.render_xlsx(["produk", "total"], [["Kopi", 7], ["Teh", None]]))).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [["produk", "total"], ["Kopi", 7], ["Teh", None]]
    assert sheet["A1"].font.bold


def _message(steps: list[dict], content: str = "Jawaban.") -> SimpleNamespace:
    return SimpleNamespace(steps_json=steps, content=content, created_at=datetime(2026, 10, 12, tzinfo=timezone.utc))


def test_export_of_a_written_document_is_the_document_itself():
    message = _message([{"type": "document", "title": "Energi Terbarukan di Indonesia", "markdown": PAPER, "words": 90}])
    pdf, name = build_export("pdf", "buatkan makalah", message, "Chat")
    text = _flat(PdfReader(io.BytesIO(pdf)))
    assert name == "energi-terbarukan-di-indonesia" and "garis khatulistiwa" in text
    assert "Laporan Analisis Data" not in text and "Pertanyaan" not in text  # not the chat-report wrapper

    docx, _ = build_export("docx", "buatkan makalah", message, "Chat")
    assert "garis khatulistiwa" in "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)


def test_export_of_an_analysis_answer():
    table = {"type": "table", "columns": ["produk", "total"], "types": ["VARCHAR", "BIGINT"], "rows": [["Kopi", 7]], "truncated": False}
    message = _message([{"type": "intent", "value": "sql"}, table], "Kopi paling laku.")
    xlsx, _ = build_export("xlsx", "apa yang laku", message, "penjualan.csv")
    assert load_workbook(io.BytesIO(xlsx)).active["A2"].value == "Kopi"
    docx, _ = build_export("docx", "apa yang laku", message, "penjualan.csv")
    document = Document(io.BytesIO(docx))
    assert "Kopi paling laku." in [p.text for p in document.paragraphs] and document.tables[0].rows[1].cells[0].text == "Kopi"
    with pytest.raises(HTTPException):
        build_export("xlsx", "halo", _message([]), "Chat")


@pytest.mark.parametrize(
    ("question", "has_table", "expected"),
    [
        ("buatkan makalah tentang energi terbarukan", False, True),
        ("tolong tulis surat lamaran kerja", False, True),
        ("susun proposal kerja sama dengan vendor laundry", True, True),
        ("buat laporan kegiatan bulan ini", False, True),
        ("buat laporan penjualan per bulan", True, False),  # with a table attached this is an analysis
        ("buatkan grafik untuk makalah saya", False, False),
        # The request that was once answered as small talk because it mentions "grafik".
        ("buatkan saya makalah untuk masalah kebakaran yang sedang viral di kahutla dan berikan data dan grafik yang bisa kamu berikan ke saya, kemudian generate ke word dan pdf", False, True),
        ("buatkan makalah karhutla lengkap dengan tabel dan diagram", True, True),
        ("apa isi makalah ini", False, False),
        ("produk apa yang paling laku", True, False),
    ],
)
def test_writing_requests_are_recognised(question, has_table, expected):
    assert is_writing_request(question, has_table) is expected


def test_strip_fence():
    assert strip_fence("```markdown\n# Judul\n\nIsi.\n```") == "# Judul\n\nIsi."
    assert strip_fence("# Judul\n\n```sql\nSELECT 1\n```\n\nIsi.") == "# Judul\n\n```sql\nSELECT 1\n```\n\nIsi."


class Writer:
    def __init__(self, hosted: bool):
        self.has_strong_model = hosted
        self.calls = []

    async def complete(self, kind, system, messages, max_tokens=4096):
        self.calls.append((system, messages[-1]["content"], max_tokens))
        return LLMResult(PAPER, "fake", "fake", 10, 500)


async def test_compose_writes_a_full_document_and_offers_it_for_download():
    llm = Writer(hosted=True)
    agent = build_agent(llm, Settings())
    state = {"question": "buatkan makalah tentang energi terbarukan", "context": "", "db_path": "", "usage": []}
    events = [e async for e, _ in stream_agent(agent, state)]

    assert [e["type"] for e in events] == ["intent", "document", "answer"]
    assert events[0]["value"] == "compose" and events[1]["title"] == "Energi Terbarukan di Indonesia"
    assert events[1]["markdown"] == PAPER.strip() and events[1]["words"] > 50
    system, _, max_tokens = llm.calls[0]
    assert "1.200 sampai 1.800 kata" in system and "Jangan mengarang kutipan" in system and max_tokens == 6000


async def test_compose_with_local_model_is_short_and_says_so():
    llm = Writer(hosted=False)
    agent = build_agent(llm, Settings())
    state = {"question": "tulis surat permohonan izin", "context": "", "db_path": "", "usage": []}
    events = [e async for e, _ in stream_agent(agent, state)]
    assert "versi ringkas dari model lokal" in events[-1]["text"]
    assert "250 sampai 350 kata" in llm.calls[0][0] and llm.calls[0][2] == 520


from app.services import chart_image  # noqa: E402

CHART_PAPER = PAPER.replace(
    "## Penutup",
    """Bauran berikut menggambarkan komposisinya.

```chart
{"type": "donut", "title": "Bauran energi listrik", "unit": "%", "source": "perkiraan", "labels": ["Batu bara", "Gas", "Terbarukan", "Lainnya"], "values": [60, 18, 14, 8]}
```

```chart
{"type": "bar", "title": "Kapasitas terpasang per sumber", "unit": "MW", "source": "dokumen", "labels": ["Surya", "Angin", "Panas bumi"], "values": [320, 150, 2400]}
```

```chart
{"type": "bar", "labels": ["A"], "values": ["bukan angka"]}
```

## Penutup""",
)


def test_chart_blocks_are_validated_and_marked_with_their_provenance():
    charts = [b.chart for b in exports.parse_markdown(CHART_PAPER) if b.kind == "chart"]
    assert [(c.kind, len(c.values)) for c in charts] == [("donut", 4), ("bar", 3)]  # the malformed third block is dropped
    assert charts[0].caption == "Angka perkiraan, perlu diverifikasi"
    assert charts[1].caption == "Sumber: dokumen yang dilampirkan"


@pytest.mark.parametrize(
    "raw",
    [
        "bukan json",
        '{"type": "garis", "labels": ["a", "b"], "values": [1, 2]}',
        '{"type": "bar", "labels": ["a", "b"], "values": [1]}',
        '{"type": "bar", "labels": ["a"], "values": [1]}',
        '{"type": "bar", "labels": ["a", "b"], "values": [1, -2]}',
        '{"type": "donut", "labels": ["a", "b"], "values": [0, 0]}',
    ],
)
def test_malformed_chart_specs_are_rejected(raw):
    with pytest.raises(chart_image.ChartError):
        chart_image.parse_spec(raw)


def test_donut_with_too_many_slices_folds_the_smallest_into_other():
    spec = chart_image.parse_spec(json.dumps({"type": "donut", "labels": list("ABCDEFGH"), "values": [30, 25, 15, 10, 8, 6, 4, 2]}))
    assert spec.labels == ["A", "B", "C", "D", "E", "Lainnya"] and spec.values[-1] == 12 and sum(spec.values) == 100


def test_charts_are_drawn_and_embedded_in_both_files():
    from PIL import Image

    for spec in (b.chart for b in exports.parse_markdown(CHART_PAPER) if b.kind == "chart"):
        png, ratio = chart_image.render_png(spec)
        image = Image.open(io.BytesIO(png))
        assert image.width == 1200 and ratio == pytest.approx(image.height / image.width)
        assert len(image.getcolors(maxcolors=100000)) > 20  # not a blank canvas

    document = Document(io.BytesIO(exports.render_docx(CHART_PAPER)))
    assert len(document.inline_shapes) == 2
    assert "```" not in "\n".join(p.text for p in document.paragraphs)

    reader = PdfReader(io.BytesIO(exports.render_pdf(CHART_PAPER)))
    assert sum(len(page.images) for page in reader.pages) == 2
    assert "chart" not in _flat(reader) and "labels" not in _flat(reader)  # the data block is drawn, not printed


async def test_compose_prompt_asks_for_tables_and_charts_with_honest_sources():
    llm = Writer(hosted=True)
    agent = build_agent(llm, Settings())
    state = {"question": "buatkan makalah tentang energi terbarukan", "context": "", "db_path": "", "usage": []}
    [e async for e, _ in stream_agent(agent, state)]
    system = llm.calls[0][0]
    assert "```chart" in system and '"donut" hanya untuk komposisi' in system
    assert "angka yang tidak ditemukan di sumber otomatis dilabeli perkiraan" in system


def test_bars_are_sorted_by_size_unless_the_labels_are_a_sequence():
    sectors = chart_image.parse_spec(json.dumps({"type": "bar", "unit": "triliun rupiah", "labels": ["Kesehatan", "Penerbangan", "Pertanian", "Pariwisata", "Pemadaman"], "values": [8, 5, 12, 4, 6]}))
    assert sectors.labels == ["Pertanian", "Kesehatan", "Pemadaman", "Penerbangan", "Pariwisata"] and sectors.values == [12, 8, 6, 5, 4]
    assert sectors.reading == "Tertinggi: Pertanian (12 triliun rupiah). Terendah: Pariwisata (4 triliun rupiah)."

    for labels in (["2015", "2019", "2023"], ["Jan", "Feb", "Mar"], ["Kuartal 1", "Kuartal 2", "Kuartal 3"], ["Tahap 1", "Tahap 2", "Tahap 3"]):
        spec = chart_image.parse_spec(json.dumps({"type": "bar", "labels": labels, "values": [5, 9, 2]}))
        assert spec.labels == labels and spec.values == [5, 9, 2], labels

    donut = chart_image.parse_spec(json.dumps({"type": "donut", "unit": "%", "labels": ["Gas", "Batu bara", "Lainnya"], "values": [18, 60, 22]}))
    assert donut.reading == "Bagian terbesar: Batu bara (60%)."


def _paper_with_chart(source: str) -> str:
    return PAPER + '\n```chart\n{"type": "bar", "source": "%s", "labels": ["A", "B"], "values": [3, 1]}\n```\n' % source


class FixedWriter(Writer):
    def __init__(self, text: str):
        super().__init__(hosted=True)
        self.text = text

    async def complete(self, kind, system, messages, max_tokens=4096):
        self.calls.append((system, messages[-1]["content"], max_tokens))
        return LLMResult(self.text, "fake", "fake", 10, 500)


async def _compose(llm, question: str, **extra) -> str:
    agent = build_agent(llm, Settings())
    state = {"question": question, "context": "", "db_path": "", "usage": [], **extra}
    return [e async for e, _ in stream_agent(agent, state)][-1]["text"]


async def test_chat_says_when_current_data_was_asked_for_but_is_not_available():
    answer = await _compose(FixedWriter(PAPER), "buatkan makalah kasus karhutla yang sekarang sedang viral, berikan datanya secara terperinci")
    assert "pencarian web belum diaktifkan" in answer and "tanpa berita atau data terbaru" in answer and "lampirkan sumbernya" in answer

    estimated = await _compose(FixedWriter(_paper_with_chart("perkiraan")), "buatkan makalah tentang energi terbarukan")
    assert "1 grafik di dalamnya memakai angka perkiraan" in estimated

    plain = await _compose(FixedWriter(PAPER), "buatkan makalah tentang energi terbarukan")
    assert "Perlu diketahui" not in plain  # nothing estimated and nothing current asked for: no warning


async def test_no_warning_when_the_document_is_written_from_attached_sources():
    chunks = [{"page": 1, "position": 0, "text": "Luas karhutla 2023 tercatat 1,16 juta hektare.", "source": "laporan.pdf"}]
    answer = await _compose(FixedWriter(_paper_with_chart("dokumen")), "buatkan makalah karhutla terbaru dari laporan ini", doc_chunks=chunks)
    assert "Perlu diketahui" not in answer


def test_pdf_table_with_a_tall_row_starts_on_the_page_where_it_is_introduced():
    filler = "Kebakaran hutan dan lahan berulang setiap musim kemarau di Indonesia. " * 14
    wordy = "Gagal panen akibat lahan terbakar dan kabut asap yang menutupi sinar matahari selama berminggu-minggu. " * 30
    markdown = (
        f"# Judul\n\n{filler}\n\nTabel berikut merangkum kerugian.\n\n"
        f"| Sektor | Dampak | Kerugian |\n|---|---|---|\n| Pertanian | {wordy} | 12 triliun |\n| Kesehatan | Lonjakan ISPA | 8 triliun |\n"
    )
    reader = PdfReader(io.BytesIO(exports.render_pdf(markdown)))
    first = " ".join(reader.pages[0].extract_text().split())
    # The first row alone is taller than the space left; it is cut at the page break rather than moved whole.
    assert len(reader.pages) == 2
    assert "Tabel berikut merangkum kerugian." in first and "Sektor" in first and "Pertanian" in first
    assert "Kesehatan" in " ".join(reader.pages[1].extract_text().split())  # and the short first column keeps the word whole


def test_a_long_chart_title_wraps_instead_of_being_cut():
    long_title = "Indikasi Lahan Terbakar di 6 Provinsi Prioritas per 8 September 2026 (hektare)"
    font = chart_image._font(17)
    lines = chart_image._wrap(long_title, font, 544 * chart_image.SCALE)
    assert len(lines) == 2 and " ".join(lines) == long_title
    short = chart_image.parse_spec(json.dumps({"type": "bar", "title": "Pendek", "labels": ["A", "B"], "values": [1, 2]}))
    long = chart_image.parse_spec(json.dumps({"type": "bar", "title": long_title, "labels": ["A", "B"], "values": [1, 2]}))
    assert chart_image.render_png(long)[1] > chart_image.render_png(short)[1]  # taller: room for the second line


def test_pdf_table_does_not_leave_a_single_row_at_the_foot_of_a_page():
    rows = "".join(f"| {year} | {n} ha | Catatan |\n" for year, n in [(2015, 26), (2019, 16), (2023, 12), (2025, 3), (2026, 20), (2027, 9)])
    table = "| Tahun | Luas | Keterangan |\n|---|---|---|\n" + rows
    for sentences in range(28, 42):  # sweep the table across the page break
        filler = "Kebakaran hutan dan lahan berulang setiap musim kemarau di Indonesia. " * sentences
        reader = PdfReader(io.BytesIO(exports.render_pdf(f"# Judul\n\n{filler}\n\n{table}")))
        start = next(page.extract_text() for page in reader.pages if "Keterangan" in page.extract_text())
        assert start.count("Catatan") >= 3, sentences  # the page where the table starts holds at least three of its rows
